#!/usr/bin/env python3
"""Fetch the public UWH Nationals draw workbook and write docs/data.json.

Usage:
    python scripts/fetch_data.py [--source PATH_OR_URL]

Default source is the public Google Sheets xlsx export. Only the `Inputs`
sheet is parsed -- it is the master draw. openpyxl is the only third-party
dependency.

Prints CHANGED (file rewritten) or UNCHANGED (payload identical apart from
`generated_at`). Exits 1 without touching docs/data.json on download failure
or if the sanity gate trips.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
import tempfile
import time
import urllib.error
import urllib.request

import openpyxl

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------

DEFAULT_SOURCE = (
    "https://docs.google.com/spreadsheets/d/"
    "1XImVBBsZdH17-jtZaWHnEmtlOl0qUqzu/export?format=xlsx"
)

SHEET_NAME = "Inputs"
EVENT_YM = "2026-09"  # year-month of the event; day comes from the day header
TIMEZONE = "Pacific/Auckland"
TRACKED_TEAMS = ["JOB HC-BD", "SGB Howick"]

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(REPO_ROOT, "docs", "data.json")

# Court -> (time, game#, white, whiteScore, black, blackScore, stage-comment)
COURT_COLUMNS = {
    1: ("B", "C", "D", "E", "F", "G", "J"),
    2: ("L", "M", "N", "O", "P", "Q", "T"),
}

DIVISIONS = {"SGA", "SGB", "SOA", "SOB", "JGA", "JGB", "JOA", "JOB"}
GROUPS = {"SG", "SO", "JG", "JO"}

PLACEHOLDER_RE = re.compile(
    r"^(?:(SG|SO|JG|JO)\s+)?(Winner|Loser)\s+G(\d+)$", re.IGNORECASE
)
DAY_NUM_RE = re.compile(r"(\d+)")
GROUP_PREFIX_RE = re.compile(r"^(SG|SO|JG|JO)\b")
# A cell of dashes is a blank placeholder, not a team (the Final Standings
# table below the draw uses '-' for every not-yet-known placing).
NOT_A_TEAM_RE = re.compile(r"^[-‐-―\s]*$")

DOWNLOAD_ATTEMPTS = 3
BACKOFF_SECONDS = [5, 15]
MIN_GAMES = 50


# --------------------------------------------------------------------------
# Source acquisition
# --------------------------------------------------------------------------


def _looks_like_url(source: str) -> bool:
    return source.startswith("http://") or source.startswith("https://")


def acquire_workbook(source: str) -> str:
    """Return a local path to the workbook, downloading it if needed.

    Exits 1 (without touching docs/data.json) if the workbook cannot be
    obtained or does not look like an xlsx file.
    """
    if not _looks_like_url(source):
        path = os.path.abspath(source)
        try:
            with open(path, "rb") as fh:
                magic = fh.read(2)
        except OSError as exc:
            sys.stderr.write("ERROR: cannot read source %s: %s\n" % (path, exc))
            sys.exit(1)
        if magic != b"PK":
            sys.stderr.write(
                "ERROR: %s does not look like an xlsx file (no PK magic)\n" % path
            )
            sys.exit(1)
        return path

    last_error = None
    for attempt in range(DOWNLOAD_ATTEMPTS):
        if attempt:
            delay = BACKOFF_SECONDS[min(attempt - 1, len(BACKOFF_SECONDS) - 1)]
            sys.stderr.write("retrying in %ss ...\n" % delay)
            time.sleep(delay)
        try:
            req = urllib.request.Request(
                source, headers={"User-Agent": "uwh-nationals-fetch/1.0"}
            )
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = resp.read()
        except (urllib.error.URLError, OSError, ValueError) as exc:
            last_error = exc
            sys.stderr.write("download attempt %d failed: %s\n" % (attempt + 1, exc))
            continue

        if not data.startswith(b"PK"):
            last_error = "response is not an xlsx file (no PK magic, %d bytes)" % len(
                data
            )
            sys.stderr.write("download attempt %d failed: %s\n" % (attempt + 1, last_error))
            continue

        fd, tmp_path = tempfile.mkstemp(prefix="uwh-draw-", suffix=".xlsx")
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        return tmp_path

    sys.stderr.write(
        "ERROR: could not download workbook after %d attempts: %s\n"
        % (DOWNLOAD_ATTEMPTS, last_error)
    )
    sys.exit(1)


# --------------------------------------------------------------------------
# Cell helpers
# --------------------------------------------------------------------------


def cell(ws, col: str, row: int):
    return ws["%s%d" % (col, row)].value


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def as_game_number(value):
    """Game number: int, float or numeric string -> int; else None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        if float(value).is_integer():
            return int(value)
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        return int(number) if number.is_integer() else None
    return None


def as_score(value):
    """Numeric -> int. '-', blank or anything else -> None (not played)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not float(value).is_integer():
            return int(round(value))
        return int(value)
    if isinstance(value, str):
        text = value.strip()
        if not text or text == "-":
            return None
        try:
            return int(round(float(text)))
        except ValueError:
            return None
    return None


def as_time(value):
    """Normalise a time cell to 'HH:MM' (NZ local); None when unusable."""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.strftime("%H:%M")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, dt.timedelta):
        total = int(value.total_seconds()) % 86400
        return "%02d:%02d" % (total // 3600, (total % 3600) // 60)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel serial time: fractional part is the fraction of a day.
        fraction = float(value) % 1.0
        minutes = int(round(fraction * 24 * 60))
        return "%02d:%02d" % ((minutes // 60) % 24, minutes % 60)
    if isinstance(value, str):
        text = value.strip()
        match = re.match(r"^(\d{1,2}):(\d{2})(?::\d{2})?", text)
        if match:
            return "%02d:%02d" % (int(match.group(1)), int(match.group(2)))
    return None


def as_version(value):
    """Draw version as a display-faithful string; None when missing."""
    if value is None:
        return None
    if isinstance(value, float) and float(value).is_integer():
        return str(int(value))
    text = as_text(value)
    return text or None


# --------------------------------------------------------------------------
# Division
# --------------------------------------------------------------------------


def is_team_name(value) -> bool:
    """A team cell must be a non-empty string that is not just dashes."""
    if not isinstance(value, str):
        return False
    text = value.strip()
    return bool(text) and not NOT_A_TEAM_RE.match(text)


def first_token(name: str) -> str:
    return name.split()[0] if name else ""


def stage_group(stage) -> str | None:
    if not stage:
        return None
    match = GROUP_PREFIX_RE.match(stage)
    if match:
        return match.group(1)
    return None


def derive_division(white: str, black: str, stage) -> str | None:
    """Division of a game.

    First token when the two sides agree on an in-set division (round robin);
    otherwise the 2-letter group taken from the stage comment prefix (cross-
    group knockouts and placeholder slots); otherwise a lone in-set token;
    otherwise null.
    """
    wt, bt = first_token(white), first_token(black)
    if wt in DIVISIONS and wt == bt:
        return wt
    group = stage_group(stage)
    if group:
        return group
    if wt in DIVISIONS:
        return wt
    if bt in DIVISIONS:
        return bt
    return None


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------


def is_day_header(ws, row: int) -> bool:
    return as_text(cell(ws, "C", row)) == "#" and bool(as_text(cell(ws, "B", row)))


def day_to_date(day_label: str) -> str | None:
    match = DAY_NUM_RE.search(day_label)
    if not match:
        return None
    return "%s-%02d" % (EVENT_YM, int(match.group(1)))


def parse_games(ws) -> list[dict]:
    games: list[dict] = []
    current_day = None
    current_date = None

    for row in range(1, ws.max_row + 1):
        if is_day_header(ws, row):
            label = as_text(cell(ws, "B", row))
            date = day_to_date(label)
            if date:
                current_day, current_date = label, date
            continue

        if current_day is None:
            continue

        for court, (c_time, c_num, c_white, c_ws, c_black, c_bs, c_stage) in sorted(
            COURT_COLUMNS.items()
        ):
            number = as_game_number(cell(ws, c_num, row))
            if number is None:
                continue
            white_cell = cell(ws, c_white, row)
            black_cell = cell(ws, c_black, row)
            if not is_team_name(white_cell) or not is_team_name(black_cell):
                continue
            # A real fixture always has a start time; the Final Standings
            # side table (place numbers + team names, no time) does not.
            game_time = as_time(cell(ws, c_time, row))
            if game_time is None:
                continue
            white = as_text(white_cell)
            black = as_text(black_cell)

            white_score = as_score(cell(ws, c_ws, row))
            black_score = as_score(cell(ws, c_bs, row))
            stage = as_text(cell(ws, c_stage, row)) or None

            games.append(
                {
                    "game": number,
                    "day": current_day,
                    "date": current_date,
                    "time": game_time,
                    "court": court,
                    "white": white,
                    "black": black,
                    "white_score": white_score,
                    "black_score": black_score,
                    "played": white_score is not None and black_score is not None,
                    "stage": stage,
                    "division": derive_division(white, black, stage),
                    "white_source": None,
                    "black_source": None,
                }
            )

    games.sort(key=lambda g: (g["date"] or "", g["time"] or "", g["court"]))
    return games


# --------------------------------------------------------------------------
# Knockout placeholder resolution
# --------------------------------------------------------------------------


def resolve_placeholders(games: list[dict]) -> None:
    by_number = {g["game"]: g for g in games}
    originals = {
        (g["game"], side): g[side] for g in games for side in ("white", "black")
    }

    for _ in range(5):
        changed = False
        for game in games:
            for side in ("white", "black"):
                match = PLACEHOLDER_RE.match(game[side])
                if not match:
                    continue
                ref = by_number.get(int(match.group(3)))
                if ref is None or not ref["played"]:
                    continue
                if ref["white_score"] == ref["black_score"]:
                    continue  # a draw resolves to nothing
                white_won = ref["white_score"] > ref["black_score"]
                want_winner = match.group(2).lower() == "winner"
                resolved = (
                    ref["white"] if white_won == want_winner else ref["black"]
                )
                if resolved and resolved != game[side]:
                    game[side] = resolved
                    changed = True
        if not changed:
            break

    for game in games:
        for side in ("white", "black"):
            original = originals[(game["game"], side)]
            if not PLACEHOLDER_RE.match(original):
                continue
            if PLACEHOLDER_RE.match(game[side]):
                # Never fully resolved: keep the sheet's own placeholder.
                game[side] = original
                game["%s_source" % side] = None
            else:
                game["%s_source" % side] = original


# --------------------------------------------------------------------------
# Payload
# --------------------------------------------------------------------------


def build_payload(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if SHEET_NAME not in wb.sheetnames:
        sys.stderr.write("ERROR: workbook has no '%s' sheet\n" % SHEET_NAME)
        sys.exit(1)
    ws = wb[SHEET_NAME]

    games = parse_games(ws)
    resolve_placeholders(games)

    return {
        "tournament": as_text(cell(ws, "B", 2)),
        "venue": as_text(cell(ws, "B", 3)),
        "draw_version": as_version(cell(ws, "Q", 4)),
        "generated_at": dt.datetime.now(dt.timezone.utc).strftime(
            "%Y-%m-%dT%H:%M:%SZ"
        ),
        "timezone": TIMEZONE,
        "tracked_teams": list(TRACKED_TEAMS),
        "games": games,
    }


def load_existing(path: str):
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return None


def without_generated_at(payload):
    return {k: v for k, v in payload.items() if k != "generated_at"}


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="Fetch the UWH Nationals draw and write docs/data.json"
    )
    parser.add_argument(
        "--source",
        default=DEFAULT_SOURCE,
        help="workbook URL or local path (default: the public Google export)",
    )
    args = parser.parse_args(argv)

    workbook_path = acquire_workbook(args.source)
    payload = build_payload(workbook_path)
    existing = load_existing(OUT_PATH)

    # Sanity gate: never trade a good file for a suspiciously small parse.
    if len(payload["games"]) < MIN_GAMES and existing is not None:
        existing_count = len(existing.get("games", []))
        if existing_count > len(payload["games"]):
            sys.stderr.write(
                "ERROR: refusing to overwrite docs/data.json — parsed only %d games "
                "(< %d) while the existing file has %d. Upstream sheet looks broken.\n"
                % (len(payload["games"]), MIN_GAMES, existing_count)
            )
            return 1

    if existing is not None and without_generated_at(existing) == without_generated_at(
        payload
    ):
        print("UNCHANGED")
        return 0

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    print("CHANGED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
